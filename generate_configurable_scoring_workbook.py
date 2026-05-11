#!/usr/bin/env python
"""Create the configurable scoring workbook used by the MVP dashboard tab."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from extras_discipline_dashboard import (
    build_scorecard_batting_innings_stats,
    build_scorecard_bowling_innings_stats,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the configurable scoring workbook.")
    parser.add_argument(
        "--scorecards-dir",
        type=Path,
        default=Path("scorecards"),
        help="Directory containing scorecard JSON files",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("assets") / "configurable_scoring.xlsx",
        help="Output workbook path",
    )
    parser.add_argument("--batting-runs-weight", type=float, default=0.75)
    parser.add_argument("--batting-delta-sr-weight", type=float, default=0.25)
    parser.add_argument("--batting-min-balls", type=int, default=20)
    parser.add_argument("--bowling-wickets-weight", type=float, default=0.6)
    parser.add_argument("--bowling-economy-weight", type=float, default=0.2)
    parser.add_argument("--bowling-sr-weight", type=float, default=0.2)
    parser.add_argument("--bowling-avg-weight", type=float, default=0)
    parser.add_argument("--bowling-min-balls", type=int, default=30)
    return parser.parse_args()


def export_configurable_scoring_workbook(
    scorecards_dir: Path,
    output_path: Path,
    *,
    default_batting_option_2: dict[str, float],
    default_bowling_option_2: dict[str, float],
    default_minimum_balls_faced: int = 20,
    default_minimum_balls_bowled: int = 30,
    default_bowling_sr_weight: float = 0.2,
    default_bowling_avg_weight: float = 0,
) -> Path:
    batting_df = build_scorecard_batting_innings_stats(scorecards_dir).copy()
    bowling_df = build_scorecard_bowling_innings_stats(scorecards_dir).copy()

    workbook = Workbook()
    inputs_ws = workbook.active
    inputs_ws.title = "Inputs"
    inputs_ws["A1"] = "Configurable scoring inputs"
    inputs_ws["A1"].font = Font(bold=True)
    input_rows = [
        ("Batting option 2 runs weight", default_batting_option_2["runs"]),
        ("Batting option 2 delta SR weight", default_batting_option_2["delta_sr"]),
        ("Batting option 2 total minimum balls faced", default_minimum_balls_faced),
        ("Bowling option 2 wickets weight", default_bowling_option_2["wickets"]),
        ("Bowling option 2 economy weight", default_bowling_option_2["economy"]),
        ("Bowling option 2 SR weight", default_bowling_sr_weight),
        ("Bowling option 2 average weight", default_bowling_avg_weight),
        ("Bowling option 2 total minimum balls bowled", default_minimum_balls_bowled),
    ]
    for row_index, (label, value) in enumerate(input_rows, start=2):
        inputs_ws[f"A{row_index}"] = label
        inputs_ws[f"B{row_index}"] = value
    inputs_ws["A10"] = "Batting option 2 equation"
    inputs_ws["A10"].font = Font(bold=True)
    inputs_ws["A11"] = r"$$V_{bat,2} = w_r R^* + \mathbf{1}[B_{tot} \ge B_{\min}] w_{sr} \Delta SR^*$$"
    inputs_ws["A13"] = "Bowling option 2 equation"
    inputs_ws["A13"].font = Font(bold=True)
    inputs_ws["A14"] = r"$$V_{bowl,2} = w_w W^* + \mathbf{1}[LB_{tot} \ge LB_{\min}] (w_e ECO^*_{inv} + w_{sr} SR^*_{inv} + w_{avg} AVG^*_{inv})$$"

    batting_ws = workbook.create_sheet("BattingInnings")
    batting_headers = [
        "match_key",
        "match_label",
        "batting_team",
        "batter",
        "runs",
        "balls",
        "strike_rate",
        "team_tournament_runs",
        "team_tournament_balls",
        "team_tournament_strike_rate",
        "delta_strike_rate",
        "runs_norm",
        "delta_sr_norm",
        "batting_option_2_score",
    ]
    batting_ws.append(batting_headers)
    batting_rows_start = 2
    batting_rows_end = batting_rows_start + len(batting_df) - 1
    for _, row in batting_df.iterrows():
        batting_ws.append(
            [
                row.get("match_key", ""),
                row.get("match_label", ""),
                row.get("batting_team", ""),
                row.get("batsman", ""),
                row.get("runs", 0),
                row.get("balls", 0),
            ]
        )
    if batting_rows_end >= batting_rows_start:
        for row_index in range(batting_rows_start, batting_rows_end + 1):
            batting_ws[f"G{row_index}"] = f'=IF(F{row_index}=0,0,E{row_index}/F{row_index}*100)'
            batting_ws[f"H{row_index}"] = f'=SUMIF($C${batting_rows_start}:$C${batting_rows_end},C{row_index},$E${batting_rows_start}:$E${batting_rows_end})'
            batting_ws[f"I{row_index}"] = f'=SUMIF($C${batting_rows_start}:$C${batting_rows_end},C{row_index},$F${batting_rows_start}:$F${batting_rows_end})'
            batting_ws[f"J{row_index}"] = f'=IF(I{row_index}=0,0,H{row_index}/I{row_index}*100)'
            batting_ws[f"K{row_index}"] = f'=MAX(0,G{row_index}-J{row_index})'
            batting_ws[f"L{row_index}"] = (
                f'=IF(MAX($E${batting_rows_start}:$E${batting_rows_end})=MIN($E${batting_rows_start}:$E${batting_rows_end}),'
                f'IF(E{row_index}>0,1,0),(E{row_index}-MIN($E${batting_rows_start}:$E${batting_rows_end}))/'
                f'(MAX($E${batting_rows_start}:$E${batting_rows_end})-MIN($E${batting_rows_start}:$E${batting_rows_end})))'
            )
            batting_ws[f"M{row_index}"] = (
                f'=IF(MAX($K${batting_rows_start}:$K${batting_rows_end})=MIN($K${batting_rows_start}:$K${batting_rows_end}),'
                f'IF(K{row_index}>0,1,0),(K{row_index}-MIN($K${batting_rows_start}:$K${batting_rows_end}))/'
                f'(MAX($K${batting_rows_start}:$K${batting_rows_end})-MIN($K${batting_rows_start}:$K${batting_rows_end})))'
            )
            batting_ws[f"N{row_index}"] = f'=Inputs!$B$2*L{row_index}+Inputs!$B$3*M{row_index}'

    bowling_ws = workbook.create_sheet("BowlingInnings")
    bowling_headers = [
        "match_key",
        "match_label",
        "bowling_team",
        "bowler",
        "wickets",
        "legal_balls",
        "overs_bowled",
        "runs_conceded",
        "economy",
        "bowling_strike_rate",
        "bowling_average",
        "wickets_norm",
        "economy_inv_norm",
        "sr_inv_norm",
        "avg_inv_norm",
        "bowling_option_2_score",
    ]
    bowling_ws.append(bowling_headers)
    bowling_rows_start = 2
    bowling_rows_end = bowling_rows_start + len(bowling_df) - 1
    for _, row in bowling_df.iterrows():
        bowling_ws.append(
            [
                row.get("match_key", ""),
                row.get("match_label", ""),
                row.get("bowling_team", ""),
                row.get("bowler", ""),
                row.get("wickets", 0),
                row.get("legal_balls", 0),
                row.get("overs_bowled", 0.0),
                row.get("runs_conceded", 0),
            ]
        )
    if bowling_rows_end >= bowling_rows_start:
        for row_index in range(bowling_rows_start, bowling_rows_end + 1):
            bowling_ws[f"I{row_index}"] = f'=IF(G{row_index}=0,0,H{row_index}/G{row_index})'
            bowling_ws[f"J{row_index}"] = f'=IF(E{row_index}=0,0,F{row_index}/E{row_index})'
            bowling_ws[f"K{row_index}"] = f'=IF(E{row_index}=0,0,H{row_index}/E{row_index})'
            bowling_ws[f"L{row_index}"] = (
                f'=IF(MAX($E${bowling_rows_start}:$E${bowling_rows_end})=MIN($E${bowling_rows_start}:$E${bowling_rows_end}),'
                f'IF(E{row_index}>0,1,0),(E{row_index}-MIN($E${bowling_rows_start}:$E${bowling_rows_end}))/'
                f'(MAX($E${bowling_rows_start}:$E${bowling_rows_end})-MIN($E${bowling_rows_start}:$E${bowling_rows_end})))'
            )
            bowling_ws[f"M{row_index}"] = (
                f'=IF(MAX($I${bowling_rows_start}:$I${bowling_rows_end})=MIN($I${bowling_rows_start}:$I${bowling_rows_end}),'
                f'IF(I{row_index}>=0,1,0),1-((I{row_index}-MIN($I${bowling_rows_start}:$I${bowling_rows_end}))/'
                f'(MAX($I${bowling_rows_start}:$I${bowling_rows_end})-MIN($I${bowling_rows_start}:$I${bowling_rows_end}))))'
            )
            bowling_ws[f"N{row_index}"] = (
                f'=IF(MAX($J${bowling_rows_start}:$J${bowling_rows_end})=MIN($J${bowling_rows_start}:$J${bowling_rows_end}),'
                f'IF(J{row_index}>=0,1,0),1-((J{row_index}-MIN($J${bowling_rows_start}:$J${bowling_rows_end}))/'
                f'(MAX($J${bowling_rows_start}:$J${bowling_rows_end})-MIN($J${bowling_rows_start}:$J${bowling_rows_end}))))'
            )
            bowling_ws[f"O{row_index}"] = (
                f'=IF(MAX($K${bowling_rows_start}:$K${bowling_rows_end})=MIN($K${bowling_rows_start}:$K${bowling_rows_end}),'
                f'IF(K{row_index}>=0,1,0),1-((K{row_index}-MIN($K${bowling_rows_start}:$K${bowling_rows_end}))/'
                f'(MAX($K${bowling_rows_start}:$K${bowling_rows_end})-MIN($K${bowling_rows_start}:$K${bowling_rows_end}))))'
            )
            bowling_ws[f"P{row_index}"] = f'=Inputs!$B$5*L{row_index}+Inputs!$B$6*M{row_index}+Inputs!$B$7*N{row_index}+Inputs!$B$8*O{row_index}'

    batting_summary_ws = workbook.create_sheet("BattingSummary")
    batting_summary_headers = [
        "team",
        "player",
        "innings_count",
        "total_score",
        "average_score",
        "batting_runs_component",
        "batting_delta_sr_component",
        "batting_total_balls_gate",
        "total_runs",
        "total_balls",
        "player_total_strike_rate",
        "tournament_total_runs",
        "tournament_total_balls",
        "tournament_total_strike_rate",
        "total_delta_strike_rate",
        "runs_norm",
        "delta_sr_norm",
    ]
    batting_summary_ws.append(batting_summary_headers)
    batting_players = (
        batting_df[["batting_team", "batsman"]].drop_duplicates().sort_values(["batting_team", "batsman"])
        if not batting_df.empty
        else batting_df.reindex(columns=["batting_team", "batsman"])
    )
    batting_summary_start = 2
    for _, row in batting_players.iterrows():
        batting_summary_ws.append([row["batting_team"], row["batsman"]])
    batting_summary_end = batting_summary_start + len(batting_players) - 1
    if batting_summary_end >= batting_summary_start and batting_rows_end >= batting_rows_start:
        for row_index in range(batting_summary_start, batting_summary_end + 1):
            batting_summary_ws[f"C{row_index}"] = (
                f'=COUNTIFS(BattingInnings!$C$2:$C${batting_rows_end},A{row_index},BattingInnings!$D$2:$D${batting_rows_end},B{row_index})'
            )
            batting_summary_ws[f"I{row_index}"] = (
                f'=SUMIFS(BattingInnings!$E$2:$E${batting_rows_end},BattingInnings!$C$2:$C${batting_rows_end},A{row_index},'
                f'BattingInnings!$D$2:$D${batting_rows_end},B{row_index})'
            )
            batting_summary_ws[f"J{row_index}"] = (
                f'=SUMIFS(BattingInnings!$F$2:$F${batting_rows_end},BattingInnings!$C$2:$C${batting_rows_end},A{row_index},'
                f'BattingInnings!$D$2:$D${batting_rows_end},B{row_index})'
            )
            batting_summary_ws[f"K{row_index}"] = f'=IF(J{row_index}=0,0,I{row_index}/J{row_index}*100)'
            batting_summary_ws[f"L{row_index}"] = (
                f'=SUM(BattingInnings!$E$2:$E${batting_rows_end})'
            )
            batting_summary_ws[f"M{row_index}"] = (
                f'=SUM(BattingInnings!$F$2:$F${batting_rows_end})'
            )
            batting_summary_ws[f"N{row_index}"] = f'=IF(M{row_index}=0,0,L{row_index}/M{row_index}*100)'
            batting_summary_ws[f"O{row_index}"] = f'=MAX(0,K{row_index}-N{row_index})'
            batting_summary_ws[f"P{row_index}"] = (
                f'=IF(MAX($I${batting_summary_start}:$I${batting_summary_end})=MIN($I${batting_summary_start}:$I${batting_summary_end}),'
                f'IF(I{row_index}>0,1,0),(I{row_index}-MIN($I${batting_summary_start}:$I${batting_summary_end}))/'
                f'(MAX($I${batting_summary_start}:$I${batting_summary_end})-MIN($I${batting_summary_start}:$I${batting_summary_end})))'
            )
            batting_summary_ws[f"Q{row_index}"] = (
                f'=IF(MAX($O${batting_summary_start}:$O${batting_summary_end})=MIN($O${batting_summary_start}:$O${batting_summary_end}),'
                f'IF(O{row_index}>0,1,0),(O{row_index}-MIN($O${batting_summary_start}:$O${batting_summary_end}))/'
                f'(MAX($O${batting_summary_start}:$O${batting_summary_end})-MIN($O${batting_summary_start}:$O${batting_summary_end})))'
            )
            batting_summary_ws[f"H{row_index}"] = f'=IF(J{row_index}>=Inputs!$B$4,1,0)'
            batting_summary_ws[f"F{row_index}"] = (
                f'=Inputs!$B$2*P{row_index}'
            )
            batting_summary_ws[f"G{row_index}"] = (
                f'=H{row_index}*Inputs!$B$3*Q{row_index}'
            )
            batting_summary_ws[f"D{row_index}"] = f'=F{row_index}+G{row_index}'
            batting_summary_ws[f"E{row_index}"] = f'=D{row_index}'

    bowling_summary_ws = workbook.create_sheet("BowlingSummary")
    bowling_summary_headers = [
        "team",
        "player",
        "innings_count",
        "total_score",
        "average_score",
        "bowling_wickets_component",
        "bowling_total_balls_gate",
        "bowling_economy_component",
        "total_wickets",
        "total_legal_balls",
        "total_runs_conceded",
        "total_overs_bowled",
        "total_economy",
        "total_bowling_strike_rate",
        "total_bowling_average",
        "wickets_norm",
        "economy_inv_norm",
        "sr_inv_norm",
        "avg_inv_norm",
        "bowling_sr_component",
        "bowling_avg_component",
    ]
    bowling_summary_ws.append(bowling_summary_headers)
    bowling_players = (
        bowling_df[["bowling_team", "bowler"]].drop_duplicates().sort_values(["bowling_team", "bowler"])
        if not bowling_df.empty
        else bowling_df.reindex(columns=["bowling_team", "bowler"])
    )
    bowling_summary_start = 2
    for _, row in bowling_players.iterrows():
        bowling_summary_ws.append([row["bowling_team"], row["bowler"]])
    bowling_summary_end = bowling_summary_start + len(bowling_players) - 1
    if bowling_summary_end >= bowling_summary_start and bowling_rows_end >= bowling_rows_start:
        for row_index in range(bowling_summary_start, bowling_summary_end + 1):
            bowling_summary_ws[f"C{row_index}"] = (
                f'=COUNTIFS(BowlingInnings!$C$2:$C${bowling_rows_end},A{row_index},BowlingInnings!$D$2:$D${bowling_rows_end},B{row_index})'
            )
            bowling_summary_ws[f"I{row_index}"] = (
                f'=SUMIFS(BowlingInnings!$E$2:$E${bowling_rows_end},BowlingInnings!$C$2:$C${bowling_rows_end},A{row_index},'
                f'BowlingInnings!$D$2:$D${bowling_rows_end},B{row_index})'
            )
            bowling_summary_ws[f"J{row_index}"] = (
                f'=SUMIFS(BowlingInnings!$F$2:$F${bowling_rows_end},BowlingInnings!$C$2:$C${bowling_rows_end},A{row_index},'
                f'BowlingInnings!$D$2:$D${bowling_rows_end},B{row_index})'
            )
            bowling_summary_ws[f"K{row_index}"] = (
                f'=SUMIFS(BowlingInnings!$H$2:$H${bowling_rows_end},BowlingInnings!$C$2:$C${bowling_rows_end},A{row_index},'
                f'BowlingInnings!$D$2:$D${bowling_rows_end},B{row_index})'
            )
            bowling_summary_ws[f"L{row_index}"] = f'=J{row_index}/6'
            bowling_summary_ws[f"M{row_index}"] = f'=IF(L{row_index}=0,0,K{row_index}/L{row_index})'
            bowling_summary_ws[f"N{row_index}"] = f'=IF(I{row_index}=0,0,J{row_index}/I{row_index})'
            bowling_summary_ws[f"O{row_index}"] = f'=IF(I{row_index}=0,0,K{row_index}/I{row_index})'
            bowling_summary_ws[f"P{row_index}"] = (
                f'=IF(MAX($I${bowling_summary_start}:$I${bowling_summary_end})=MIN($I${bowling_summary_start}:$I${bowling_summary_end}),'
                f'IF(I{row_index}>0,1,0),(I{row_index}-MIN($I${bowling_summary_start}:$I${bowling_summary_end}))/'
                f'(MAX($I${bowling_summary_start}:$I${bowling_summary_end})-MIN($I${bowling_summary_start}:$I${bowling_summary_end})))'
            )
            bowling_summary_ws[f"Q{row_index}"] = (
                f'=IF(MAX($M${bowling_summary_start}:$M${bowling_summary_end})=MIN($M${bowling_summary_start}:$M${bowling_summary_end}),'
                f'IF(M{row_index}>=0,1,0),1-((M{row_index}-MIN($M${bowling_summary_start}:$M${bowling_summary_end}))/'
                f'(MAX($M${bowling_summary_start}:$M${bowling_summary_end})-MIN($M${bowling_summary_start}:$M${bowling_summary_end}))))'
            )
            bowling_summary_ws[f"R{row_index}"] = (
                f'=IF(MAX($N${bowling_summary_start}:$N${bowling_summary_end})=MIN($N${bowling_summary_start}:$N${bowling_summary_end}),'
                f'IF(N{row_index}>=0,1,0),1-((N{row_index}-MIN($N${bowling_summary_start}:$N${bowling_summary_end}))/'
                f'(MAX($N${bowling_summary_start}:$N${bowling_summary_end})-MIN($N${bowling_summary_start}:$N${bowling_summary_end}))))'
            )
            bowling_summary_ws[f"S{row_index}"] = (
                f'=IF(MAX($O${bowling_summary_start}:$O${bowling_summary_end})=MIN($O${bowling_summary_start}:$O${bowling_summary_end}),'
                f'IF(O{row_index}>=0,1,0),1-((O{row_index}-MIN($O${bowling_summary_start}:$O${bowling_summary_end}))/'
                f'(MAX($O${bowling_summary_start}:$O${bowling_summary_end})-MIN($O${bowling_summary_start}:$O${bowling_summary_end}))))'
            )
            bowling_summary_ws[f"G{row_index}"] = f'=IF(J{row_index}>=Inputs!$B$9,1,0)'
            bowling_summary_ws[f"F{row_index}"] = (
                f'=Inputs!$B$5*P{row_index}'
            )
            bowling_summary_ws[f"H{row_index}"] = (
                f'=G{row_index}*Inputs!$B$6*Q{row_index}'
            )
            bowling_summary_ws[f"T{row_index}"] = (
                f'=G{row_index}*Inputs!$B$7*R{row_index}'
            )
            bowling_summary_ws[f"U{row_index}"] = (
                f'=G{row_index}*Inputs!$B$8*S{row_index}'
            )
            bowling_summary_ws[f"D{row_index}"] = f'=F{row_index}+H{row_index}+T{row_index}+U{row_index}'
            bowling_summary_ws[f"E{row_index}"] = f'=D{row_index}'

    combined_ws = workbook.create_sheet("CombinedScores")
    combined_headers = [
        "team",
        "player",
        "batting_runs_component",
        "batting_delta_sr_component",
        "batting_total_balls_gate",
        "batting_option_2",
        "bowling_wickets_component",
        "bowling_total_balls_gate",
        "bowling_economy_component",
        "bowling_sr_component",
        "bowling_avg_component",
        "bowling_option_2",
        "average_score",
        "rank",
    ]
    combined_ws.append(combined_headers)
    combined_players = pd.concat(
        [
            batting_players.rename(columns={"batting_team": "team", "batsman": "player"}),
            bowling_players.rename(columns={"bowling_team": "team", "bowler": "player"}),
        ],
        ignore_index=True,
    ).drop_duplicates()
    combined_players = combined_players.sort_values(["team", "player"]) if not combined_players.empty else combined_players
    combined_start = 2
    for _, row in combined_players.iterrows():
        combined_ws.append([row["team"], row["player"]])
    combined_end = combined_start + len(combined_players) - 1
    if combined_end >= combined_start:
        batting_summary_range_end = max(batting_summary_end, batting_summary_start)
        bowling_summary_range_end = max(bowling_summary_end, bowling_summary_start)
        for row_index in range(combined_start, combined_end + 1):
            combined_ws[f"C{row_index}"] = (
                f'=SUMIFS(BattingSummary!$F$2:$F${batting_summary_range_end},BattingSummary!$A$2:$A${batting_summary_range_end},A{row_index},'
                f'BattingSummary!$B$2:$B${batting_summary_range_end},B{row_index})'
            )
            combined_ws[f"D{row_index}"] = (
                f'=SUMIFS(BattingSummary!$G$2:$G${batting_summary_range_end},BattingSummary!$A$2:$A${batting_summary_range_end},A{row_index},'
                f'BattingSummary!$B$2:$B${batting_summary_range_end},B{row_index})'
            )
            combined_ws[f"E{row_index}"] = (
                f'=SUMIFS(BattingSummary!$H$2:$H${batting_summary_range_end},BattingSummary!$A$2:$A${batting_summary_range_end},A{row_index},'
                f'BattingSummary!$B$2:$B${batting_summary_range_end},B{row_index})'
            )
            combined_ws[f"F{row_index}"] = (
                f'=SUMIFS(BattingSummary!$E$2:$E${batting_summary_range_end},BattingSummary!$A$2:$A${batting_summary_range_end},A{row_index},'
                f'BattingSummary!$B$2:$B${batting_summary_range_end},B{row_index})'
            )
            combined_ws[f"G{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$F$2:$F${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"H{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$G$2:$G${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"I{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$H$2:$H${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"J{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$T$2:$T${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"K{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$U$2:$U${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"L{row_index}"] = (
                f'=SUMIFS(BowlingSummary!$E$2:$E${bowling_summary_range_end},BowlingSummary!$A$2:$A${bowling_summary_range_end},A{row_index},'
                f'BowlingSummary!$B$2:$B${bowling_summary_range_end},B{row_index})'
            )
            combined_ws[f"M{row_index}"] = f'=(F{row_index}+L{row_index})/2'
            combined_ws[f"N{row_index}"] = (
                f'=RANK(M{row_index},$M${combined_start}:$M${combined_end},0)+COUNTIF($M${combined_start}:M{row_index},M{row_index})-1'
            )

    top5_ws = workbook.create_sheet("Top5")
    top5_headers = [
        "rank",
        "team",
        "player",
        "batting_runs_component",
        "batting_delta_sr_component",
        "batting_total_balls_gate",
        "batting_option_2",
        "bowling_wickets_component",
        "bowling_total_balls_gate",
        "bowling_economy_component",
        "bowling_sr_component",
        "bowling_avg_component",
        "bowling_option_2",
        "average_score",
    ]
    top5_ws.append(top5_headers)
    if combined_end >= combined_start:
        for rank_value in range(1, 6):
            row_index = rank_value + 1
            top5_ws[f"A{row_index}"] = rank_value
            for column_letter, combined_column in zip(
                ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"],
                ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"],
                strict=True,
            ):
                top5_ws[f"{column_letter}{row_index}"] = (
                    f'=IFERROR(INDEX(CombinedScores!${combined_column}$2:${combined_column}${combined_end},'
                    f'MATCH(A{row_index},CombinedScores!$N$2:$N${combined_end},0)),0)'
                )
            top5_ws[f"B{row_index}"] = (
                f'=IFERROR(INDEX(CombinedScores!$A$2:$A${combined_end},MATCH(A{row_index},CombinedScores!$N$2:$N${combined_end},0)),"")'
            )
            top5_ws[f"C{row_index}"] = (
                f'=IFERROR(INDEX(CombinedScores!$B$2:$B${combined_end},MATCH(A{row_index},CombinedScores!$N$2:$N${combined_end},0)),"")'
            )

    for worksheet in workbook.worksheets:
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)
        worksheet.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main() -> int:
    args = parse_args()
    export_configurable_scoring_workbook(
        args.scorecards_dir,
        args.output,
        default_batting_option_2={
            "runs": args.batting_runs_weight,
            "delta_sr": args.batting_delta_sr_weight,
        },
        default_bowling_option_2={
            "wickets": args.bowling_wickets_weight,
            "economy": args.bowling_economy_weight,
        },
        default_minimum_balls_faced=args.batting_min_balls,
        default_minimum_balls_bowled=args.bowling_min_balls,
        default_bowling_sr_weight=args.bowling_sr_weight,
        default_bowling_avg_weight=args.bowling_avg_weight,
    )
    print(f"{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
