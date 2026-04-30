from __future__ import annotations

import json
from pathlib import Path

import extras_discipline_dashboard
import pandas as pd
import pytest
from openpyxl import load_workbook

from generate_configurable_scoring_workbook import export_configurable_scoring_workbook
from extras_discipline_dashboard import (
    apply_scorecard_points_overrides,
    aggregate_batting_leaders,
    aggregate_bowling_leaders,
    aggregate_scorecard_batting_leaders,
    aggregate_scorecard_bowling_leaders,
    build_batting_innings_stats,
    build_bowling_innings_stats,
    build_scorecard_batting_innings_stats,
    build_scorecard_bowling_innings_stats,
    compute_batting_option_tables,
    compute_bowling_option_tables,
    compute_scorecard_all_rounder_score_table,
    build_points_table_components,
    build_top_batter_scoring_types_figure,
    build_top_bowler_wicket_types_figure,
    resolve_wicket_type,
)


def test_resolve_wicket_type_maps_commentary_subtypes() -> None:
    rows = [
        ({"dismissal_kind": "out", "commentary": "A to B, OUT Caught out, Caught by Fielder"}, "caught"),
        ({"dismissal_kind": "out", "commentary": "A to B, OUT Bowled B b A"}, "bowled"),
        ({"dismissal_kind": "out", "commentary": "A to B, OUT LBW B lbw b A"}, "lbw"),
        ({"dismissal_kind": "out", "commentary": "A to B, OUT Stumped B st Keeper b A"}, "stumped"),
        ({"dismissal_kind": "out", "commentary": "A to B, OUT Hit wicket B hit wkt b A"}, "hit_wicket"),
        ({"dismissal_kind": "run_out", "commentary": "A to B, OUT Run out"}, "run_out"),
        ({"dismissal_kind": "retired_hurt", "commentary": "A to B, retired hurt"}, "retired_hurt"),
    ]

    for row, expected in rows:
        assert resolve_wicket_type(pd.Series(row)) == expected


def test_aggregate_bowling_leaders_includes_wicket_type_counts() -> None:
    df = pd.DataFrame(
        [
            {
                "bowling_team": "Team 1",
                "bowler": "Bowler A",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "wicket_type": "caught",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
            {
                "bowling_team": "Team 1",
                "bowler": "Bowler A",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "wicket_type": "bowled",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
            {
                "bowling_team": "Team 1",
                "bowler": "Bowler A",
                "is_wicket": 1,
                "dismissal_kind": "run_out",
                "wicket_type": "run_out",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
            {
                "bowling_team": "Team 2",
                "bowler": "Bowler B",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "wicket_type": "hit_wicket",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
        ]
    )

    result = aggregate_bowling_leaders(df)

    bowler_a = result[result["bowler"] == "Bowler A"].iloc[0]
    assert bowler_a["wickets"] == 2
    assert bowler_a["caught"] == 1
    assert bowler_a["bowled"] == 1
    assert bowler_a["hit_wicket"] == 0

    bowler_b = result[result["bowler"] == "Bowler B"].iloc[0]
    assert bowler_b["wickets"] == 1
    assert bowler_b["hit_wicket"] == 1


def test_aggregate_batting_leaders_includes_scoring_shot_counts() -> None:
    df = pd.DataFrame(
        [
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 1,
                "extra_type": "none",
                "boundary_type": "",
            },
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 2,
                "extra_type": "none",
                "boundary_type": "",
            },
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 4,
                "extra_type": "none",
                "boundary_type": "FOUR",
            },
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 6,
                "extra_type": "none",
                "boundary_type": "SIX",
            },
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 1,
                "extra_type": "wide",
                "boundary_type": "",
            },
        ]
    )

    result = aggregate_batting_leaders(df)

    batter_a = result[result["batsman"] == "Batter A"].iloc[0]
    assert batter_a["runs"] == 13
    assert batter_a["balls"] == 4
    assert batter_a["ones"] == 1
    assert batter_a["twos"] == 1
    assert batter_a["fours"] == 1
    assert batter_a["sixes"] == 1
    assert batter_a["ones_runs"] == 1
    assert batter_a["twos_runs"] == 2
    assert batter_a["fours_runs"] == 4
    assert batter_a["sixes_runs"] == 6


def test_build_top_bowler_wicket_types_figure_uses_non_zero_counts() -> None:
    top_bowling_df = pd.DataFrame(
        [
            {
                "bowler": "Bowler A",
                "caught": 2,
                "bowled": 1,
                "stumped": 0,
                "hit_wicket": 1,
            }
        ]
    )

    figure = build_top_bowler_wicket_types_figure(top_bowling_df)

    trace_names = {trace.name for trace in figure.data}
    assert trace_names == {"Caught", "Bowled", "Hit wicket"}
    percentages = {trace.name: trace.y[0] for trace in figure.data}
    assert percentages == {"Caught": 50.0, "Bowled": 25.0, "Hit wicket": 25.0}
    absolute_values = {trace.name: trace.customdata[0][0] for trace in figure.data}
    assert absolute_values == {"Caught": 2, "Bowled": 1, "Hit wicket": 1}
    assert "Wickets: %{customdata[0]}" in figure.data[0].hovertemplate


def test_build_top_batter_scoring_types_figure_uses_percentage_contributions() -> None:
    top_batting_df = pd.DataFrame(
        [
            {
                "batsman": "Batter A",
                "ones_runs": 2,
                "twos_runs": 4,
                "fours_runs": 8,
                "sixes_runs": 6,
            }
        ]
    )

    figure = build_top_batter_scoring_types_figure(top_batting_df)

    trace_names = {trace.name for trace in figure.data}
    assert trace_names == {"1s", "2s", "4s", "6s"}
    percentages = {trace.name: trace.y[0] for trace in figure.data}
    assert percentages == {"1s": 10.0, "2s": 20.0, "4s": 40.0, "6s": 30.0}
    absolute_values = {trace.name: trace.customdata[0][0] for trace in figure.data}
    assert absolute_values == {"1s": 2, "2s": 4, "4s": 8, "6s": 6}
    assert "Runs: %{customdata[0]}" in figure.data[0].hovertemplate


def test_build_points_table_components_uses_grouped_bowling_headers() -> None:
    points_df = pd.DataFrame(
        [
            {
                "Team Name": "Team 1",
                "Matches": 1,
                "Won": 1,
                "Lost": 0,
                "Points": 2,
                "Net RR": 1.0,
                "For": "10/1",
                "Against": "9/2",
                "Last 5": "W",
            }
        ]
    )
    commentary_df = pd.DataFrame(
        [
            {
                "batting_team": "Team 2",
                "bowling_team": "Team 1",
                "batsman": "Batter A",
                "bowler": "Bowler A",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_boundary": 0,
                "boundary_type": "",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "wicket_type": "caught",
                "is_legal_ball": 1,
                "total_runs": 0,
            }
        ]
    )

    (
        points_table_component,
        points_table_footnote,
        batting_table_component,
        batting_table_footnote,
        bowling_table_component,
        bowling_table_footnote,
        batting_chart_figure,
        _,
    ) = build_points_table_components(points_df, commentary_df)

    assert points_table_component.columns == [
        {"name": "Team Name", "id": "Team Name"},
        {"name": "Matches", "id": "Matches"},
        {"name": "Won", "id": "Won"},
        {"name": "Lost", "id": "Lost"},
        {"name": "Points", "id": "Points"},
        {"name": "Net RR", "id": "Net RR"},
        {"name": "For", "id": "For"},
        {"name": "Against", "id": "Against"},
        {"name": "Last 5", "id": "Last 5"},
    ]
    assert points_table_component.data == [
        {
            "Team Name": "Team 1",
            "Matches": 1,
            "Won": 1,
            "Lost": 0,
            "Points": 2,
            "Net RR": 1.0,
            "For": "10/1",
            "Against": "9/2",
            "Last 5": "W",
        }
    ]
    assert points_table_footnote.children == ""
    assert batting_table_footnote.children == ""
    assert batting_table_component.merge_duplicate_headers is True
    assert batting_table_component.columns[:5] == [
        {"name": ["Identity", "Team"], "id": "batting_team"},
        {"name": ["Identity", "Batter"], "id": "batter"},
        {"name": ["Output", "Runs"], "id": "runs"},
        {"name": ["Output", "Balls"], "id": "balls"},
        {"name": ["Scoring Shots", "1s"], "id": "ones"},
    ]

    assert bowling_table_component.merge_duplicate_headers is True
    assert bowling_table_component.columns[:4] == [
        {"name": ["Identity", "Team"], "id": "bowling_team"},
        {"name": ["Identity", "Bowler"], "id": "bowler"},
        {"name": ["Wickets", "Total"], "id": "wickets"},
        {"name": ["Wicket Types", "Caught"], "id": "caught"},
    ]
    assert bowling_table_footnote.children == ""
    assert batting_chart_figure.layout.title.text == "Scoring-Type Run Contribution for Top 5 Batters"


def test_apply_scorecard_points_overrides_marks_derived_columns(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "past_matches.json").write_text(
        json.dumps(
            [
                {
                    "match_id": 101,
                    "team_a": "Team 1",
                    "team_b": "Team 2",
                    "winning_team": "Team 1",
                    "match_result": "Resulted",
                    "match_summary": {"summary": "Team 1 won by 10 runs"},
                    "match_end_time": "2026-04-01 10:00:00",
                    "match_start_time": "2026-04-01T09:00:00.000Z",
                    "overs": 13,
                    "balls": 0,
                }
            ]
        ),
        encoding="utf-8",
    )
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "innings": [
                    {
                        "teamName": "Team 1",
                        "inning": 1,
                        "total_run": 70,
                        "total_wicket": 10,
                        "overs_played": "12.2",
                    },
                    {
                        "teamName": "Team 2",
                        "inning": 2,
                        "total_run": 60,
                        "total_wicket": 8,
                        "overs_played": "13.0",
                    },
                ],
            }
        ),
        encoding="utf-8",
    )

    points_df = pd.DataFrame(
        [
            {
                "Team Name": "Team 1",
                "Matches": 0,
                "Won": 0,
                "Lost": 0,
                "Points": 0,
                "Net RR": 0.0,
                "For": "0/0",
                "Against": "0/0",
                "Last 5": "",
            }
        ]
    )

    result = apply_scorecard_points_overrides(points_df, scorecards_dir)

    row = result.iloc[0].to_dict()
    assert row["Team Name"] == "Team 1"
    assert row["Matches"] == 1
    assert row["Won"] == 1
    assert row["Lost"] == 0
    assert row["Points"] == 2
    assert round(row["Net RR"], 6) == round(0.7692307692307692, 6)
    assert row["For"] == "70/13"
    assert row["Against"] == "60/13"
    assert row["Last 5"] == "W"
    assert result.attrs["derived_points_columns"] == [
        "Against",
        "For",
        "Last 5",
        "Lost",
        "Matches",
        "Net RR",
        "Points",
        "Won",
    ]


def test_aggregate_scorecard_batting_leaders_prefers_scorecard_rows(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "scorecard": [
                    {
                        "teamName": "Team 1",
                        "batting": [
                            {"name": "Batter A", "runs": 40, "balls": 20, "4s": 3, "6s": 2},
                            {"name": "Batter B", "runs": 10, "balls": 12, "4s": 0, "6s": 0},
                        ],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    commentary_df = pd.DataFrame(
        [
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 1,
                "extra_type": "none",
                "boundary_type": "",
            },
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 2,
                "extra_type": "none",
                "boundary_type": "",
            },
        ]
    )

    result = aggregate_scorecard_batting_leaders(scorecards_dir, commentary_df)

    batter_a = result[result["batsman"] == "Batter A"].iloc[0]
    assert batter_a["runs"] == 40
    assert batter_a["balls"] == 20
    assert batter_a["fours"] == 3
    assert batter_a["sixes"] == 2
    assert batter_a["ones"] == 1
    assert batter_a["twos"] == 1


def test_aggregate_scorecard_bowling_leaders_prefers_scorecard_rows(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "match_title": "Team 1 vs Team 2",
                "scorecard": [
                    {
                        "teamName": "Team 2",
                        "batting": [
                            {"name": "Batter A", "how_to_out": "c Fielder b Bowler A"},
                            {"name": "Batter B", "how_to_out": "b Bowler A"},
                            {"name": "Batter C", "how_to_out": "not out"},
                        ],
                        "bowling": [
                            {
                                "name": "Bowler A",
                                "overs": 3,
                                "balls": 0,
                                "runs": 12,
                                "wickets": 2,
                                "wide": 1,
                                "noball": 0,
                                "extra_run": 0,
                                "bonus_run": 0,
                            }
                        ],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    commentary_df = pd.DataFrame(
        [
            {
                "batting_team": "Team 2",
                "bowling_team": "Team 1",
            }
        ]
    )

    result = aggregate_scorecard_bowling_leaders(scorecards_dir, commentary_df)

    bowler_a = result[result["bowler"] == "Bowler A"].iloc[0]
    assert bowler_a["bowling_team"] == "Team 1"
    assert bowler_a["wickets"] == 2
    assert bowler_a["caught"] == 1
    assert bowler_a["bowled"] == 1
    assert bowler_a["runs_conceded"] == 12
    assert bowler_a["extras"] == 1


def test_build_points_table_components_marks_scorecard_derived_batting_and_bowling_columns(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "match_title": "Team 1 vs Team 2",
                "scorecard": [
                    {
                        "teamName": "Team 1",
                        "batting": [{"name": "Batter A", "runs": 20, "balls": 10, "4s": 2, "6s": 1}],
                        "bowling": [],
                    },
                    {
                        "teamName": "Team 2",
                        "batting": [{"name": "Batter B", "how_to_out": "b Bowler A"}],
                        "bowling": [
                            {
                                "name": "Bowler A",
                                "overs": 2,
                                "balls": 0,
                                "runs": 10,
                                "wickets": 1,
                                "wide": 1,
                                "noball": 0,
                                "extra_run": 0,
                                "bonus_run": 0,
                            }
                        ],
                    },
                ],
            }
        ),
        encoding="utf-8",
    )
    points_df = pd.DataFrame(
        [
            {
                "Team Name": "Team 1",
                "Matches": 1,
                "Won": 1,
                "Lost": 0,
                "Points": 2,
                "Net RR": 1.0,
                "For": "20/2",
                "Against": "10/2",
                "Last 5": "W",
            }
        ]
    )
    commentary_df = pd.DataFrame(
        [
            {
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 1,
                "extra_type": "none",
                "boundary_type": "",
                "bowling_team": "Team 2",
            },
            {
                "batting_team": "Team 2",
                "batsman": "Batter B",
                "batsman_runs": 0,
                "extra_type": "none",
                "boundary_type": "",
                "bowling_team": "Team 1",
            },
        ]
    )

    (
        _,
        _,
        batting_table_component,
        batting_table_footnote,
        bowling_table_component,
        bowling_table_footnote,
        _,
        _,
    ) = build_points_table_components(points_df, commentary_df, scorecards_dir)

    assert {"name": ["Output", "Runs*"], "id": "runs"} in batting_table_component.columns
    assert {"name": ["Scoring Shots", "4s*"], "id": "fours"} in batting_table_component.columns
    assert batting_table_footnote.children != ""
    assert {"name": ["Wickets", "Total*"], "id": "wickets"} in bowling_table_component.columns
    assert {"name": ["Runs", "Conceded*"], "id": "runs_conceded"} in bowling_table_component.columns
    assert bowling_table_footnote.children != ""


def test_build_batting_innings_stats_computes_team_and_batter_rates() -> None:
    df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 4,
                "boundary_runs": 4,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 2,
                "boundary_runs": 0,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter B",
                "batsman_runs": 1,
                "boundary_runs": 0,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter B",
                "batsman_runs": 1,
                "boundary_runs": 0,
                "extra_type": "wide",
            },
        ]
    )

    result = build_batting_innings_stats(df)

    batter_a = result[result["batsman"] == "Batter A"].iloc[0]
    assert batter_a["runs"] == 6
    assert batter_a["balls"] == 2
    assert batter_a["team_runs"] == 7
    assert batter_a["team_balls"] == 3
    assert round(batter_a["team_strike_rate"], 2) == round(7 / 3 * 100.0, 2)
    assert round(batter_a["delta_strike_rate"], 2) > 0


def test_compute_batting_option_tables_ranks_by_average_score() -> None:
    df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 4,
                "boundary_runs": 4,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 4,
                "boundary_runs": 4,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter B",
                "batsman_runs": 2,
                "boundary_runs": 0,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 2",
                "batsman": "Batter C",
                "batsman_runs": 1,
                "boundary_runs": 0,
                "extra_type": "none",
            },
        ]
    )

    result = compute_batting_option_tables(
        df,
        runs_weight_option_2=0.7,
        strike_rate_weight_option_2=0.3,
        runs_weight_option_3=0.5,
        strike_rate_weight_option_3=0.25,
        boundary_weight_option_3=0.25,
    )

    assert result["option_1"].iloc[0]["batsman"] == "Batter A"
    assert result["option_2"].iloc[0]["batsman"] == "Batter A"
    assert result["option_3"].iloc[0]["batsman"] == "Batter A"


def test_compute_batting_option_tables_option_3_applies_minimum_balls_gate() -> None:
    df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 10,
                "boundary_runs": 8,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter A",
                "batsman_runs": 2,
                "boundary_runs": 0,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter B",
                "batsman_runs": 6,
                "boundary_runs": 4,
                "extra_type": "none",
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "batting_team": "Team 1",
                "batsman": "Batter C",
                "batsman_runs": 2,
                "boundary_runs": 0,
                "extra_type": "none",
            },
        ]
    )

    result = compute_batting_option_tables(
        df,
        runs_weight_option_2=0.7,
        strike_rate_weight_option_2=0.3,
        runs_weight_option_3=0.75,
        strike_rate_weight_option_3=0.25,
        boundary_weight_option_3=0.0,
        minimum_balls_faced_option_3=2,
    )

    batter_b_row = result["option_3"][result["option_3"]["batsman"] == "Batter B"].iloc[0]
    assert batter_b_row["innings_count"] == 1
    assert batter_b_row["total_score"] > 0.0
    assert batter_b_row["average_score"] > 0.0


def test_compute_bowling_option_tables_rewards_wickets_and_economy() -> None:
    df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "bowling_team": "Team 1",
                "bowler": "Bowler A",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "bowling_team": "Team 1",
                "bowler": "Bowler A",
                "is_wicket": 1,
                "dismissal_kind": "out",
                "batsman_runs": 0,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
            {
                "match_key": "m1",
                "match_label": "Match 1",
                "bowling_team": "Team 1",
                "bowler": "Bowler B",
                "is_wicket": 0,
                "dismissal_kind": "",
                "batsman_runs": 1,
                "extras": 0,
                "extra_type": "none",
                "is_legal_ball": 1,
            },
        ]
    )

    innings_stats = build_bowling_innings_stats(df)
    assert innings_stats[innings_stats["bowler"] == "Bowler A"].iloc[0]["wickets"] == 2

    result = compute_bowling_option_tables(df, wickets_weight_option_2=0.6, economy_weight_option_2=0.4)
    assert result["option_1"].iloc[0]["bowler"] == "Bowler A"
    assert result["option_2"].iloc[0]["bowler"] == "Bowler A"


def test_compute_scorecard_all_rounder_score_table_uses_scorecard_inputs(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "past_matches.json").write_text(
        json.dumps(
            [
                {
                    "match_id": 101,
                    "team_a": "Team 1",
                    "team_b": "Team 2",
                    "match_start_time": "2026-04-01T09:00:00.000Z",
                }
            ]
        ),
        encoding="utf-8",
    )
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "match_title": "Team 1 vs Team 2",
                "innings": [
                    {"teamName": "Team 1", "total_run": 30, "total_wicket": 2, "overs_played": "3.0"},
                    {"teamName": "Team 2", "total_run": 20, "total_wicket": 4, "overs_played": "3.0"},
                ],
                "scorecard": [
                    {
                        "teamName": "Team 1",
                        "batting": [
                            {"name": "Player A", "runs": 20, "balls": 10, "4s": 2, "6s": 1},
                            {"name": "Player B", "runs": 10, "balls": 8, "4s": 1, "6s": 0},
                        ],
                        "bowling": [
                            {"name": "Player C", "overs": 3, "balls": 0, "runs": 30, "wickets": 0}
                        ],
                    },
                    {
                        "teamName": "Team 2",
                        "batting": [
                            {"name": "Player C", "runs": 12, "balls": 10, "4s": 0, "6s": 1},
                            {"name": "Player D", "runs": 8, "balls": 8, "4s": 1, "6s": 0},
                        ],
                        "bowling": [
                            {"name": "Player A", "overs": 3, "balls": 0, "runs": 20, "wickets": 2},
                            {"name": "Player D", "overs": 3, "balls": 0, "runs": 10, "wickets": 1},
                        ],
                    },
                ],
            }
        ),
        encoding="utf-8",
    )

    batting_stats = build_scorecard_batting_innings_stats(scorecards_dir)
    bowling_stats = build_scorecard_bowling_innings_stats(scorecards_dir)
    assert set(batting_stats["batsman"]) >= {"Player A", "Player B", "Player C", "Player D"}
    assert set(bowling_stats["bowler"]) >= {"Player A", "Player D"}

    result = compute_scorecard_all_rounder_score_table(
        scorecards_dir,
        match_key="ALL",
        batting_runs_weight=0.5,
        batting_strike_rate_weight=0.25,
        bowling_wickets_weight=0.6,
        bowling_economy_weight=0.4,
    )

    assert list(result.columns) == [
        "team",
        "player",
        "batting_runs_component",
        "batting_delta_sr_component",
        "batting_total_balls_gate",
        "batting_option_2",
        "bowling_wickets_component",
        "bowling_total_balls_gate",
        "bowling_economy_component",
        "bowling_option_2",
        "average_score",
    ]
    assert result.iloc[0]["player"] == "Player A"
    assert result.iloc[0]["batting_runs_component"] > 0
    assert result.iloc[0]["batting_option_2"] > 0
    assert result.iloc[0]["bowling_wickets_component"] > 0
    assert result.iloc[0]["bowling_option_2"] > 0
    assert result.iloc[0]["average_score"] == (
        result.iloc[0]["batting_option_2"] + result.iloc[0]["bowling_option_2"]
    ) / 2.0


def test_compute_scorecard_all_rounder_score_table_uses_scope_totals_not_innings_means(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    batting_df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "batting_team": "Team 1",
                "batsman": "Player A",
                "runs": 30,
                "balls": 10,
            },
            {
                "match_key": "m2",
                "batting_team": "Team 1",
                "batsman": "Player A",
                "runs": 0,
                "balls": 5,
            },
            {
                "match_key": "m1",
                "batting_team": "Team 1",
                "batsman": "Player B",
                "runs": 15,
                "balls": 10,
            },
            {
                "match_key": "m1",
                "batting_team": "Team 2",
                "batsman": "Player C",
                "runs": 10,
                "balls": 10,
            },
        ]
    )
    batting_df["strike_rate"] = batting_df.apply(
        lambda row: row["runs"] / row["balls"] * 100.0 if row["balls"] > 0 else 0.0,
        axis=1,
    )
    batting_df["match_label"] = batting_df["match_key"]
    batting_df["match_title"] = batting_df["match_key"]
    team_totals = (
        batting_df.groupby("batting_team", dropna=False)
        .agg(team_tournament_runs=("runs", "sum"), team_tournament_balls=("balls", "sum"))
        .reset_index()
    )
    team_totals["team_tournament_strike_rate"] = team_totals.apply(
        lambda row: row["team_tournament_runs"] / row["team_tournament_balls"] * 100.0
        if row["team_tournament_balls"] > 0
        else 0.0,
        axis=1,
    )
    batting_df = batting_df.merge(team_totals, on="batting_team", how="left")
    batting_df["delta_strike_rate"] = (
        batting_df["strike_rate"] - batting_df["team_tournament_strike_rate"]
    ).clip(lower=0.0)

    bowling_df = pd.DataFrame(
        [
            {
                "match_key": "m1",
                "match_label": "m1",
                "match_title": "m1",
                "bowling_team": "Team 1",
                "bowler": "Player A",
                "wickets": 2,
                "legal_balls": 12,
                "overs_bowled": 2.0,
                "runs_conceded": 10,
                "economy": 5.0,
            },
            {
                "match_key": "m2",
                "match_label": "m2",
                "match_title": "m2",
                "bowling_team": "Team 1",
                "bowler": "Player A",
                "wickets": 0,
                "legal_balls": 12,
                "overs_bowled": 2.0,
                "runs_conceded": 6,
                "economy": 3.0,
            },
            {
                "match_key": "m1",
                "match_label": "m1",
                "match_title": "m1",
                "bowling_team": "Team 1",
                "bowler": "Player B",
                "wickets": 1,
                "legal_balls": 12,
                "overs_bowled": 2.0,
                "runs_conceded": 12,
                "economy": 6.0,
            },
        ]
    )

    monkeypatch.setattr(
        extras_discipline_dashboard,
        "build_scorecard_batting_innings_stats",
        lambda _scorecards_dir: batting_df.copy(),
    )
    monkeypatch.setattr(
        extras_discipline_dashboard,
        "build_scorecard_bowling_innings_stats",
        lambda _scorecards_dir: bowling_df.copy(),
    )

    result = compute_scorecard_all_rounder_score_table(
        scorecards_dir=Path("."),
        match_key="ALL",
        batting_runs_weight=0.5,
        batting_strike_rate_weight=0.25,
        batting_minimum_balls_faced=0,
        bowling_wickets_weight=0.6,
        bowling_economy_weight=0.4,
        bowling_minimum_balls_bowled=0,
    )

    player_a = result[result["player"] == "Player A"].iloc[0]
    assert player_a["batting_runs_component"] == pytest.approx(0.5)
    assert player_a["batting_delta_sr_component"] == pytest.approx(0.25)
    assert player_a["bowling_wickets_component"] == pytest.approx(0.6)
    assert player_a["bowling_total_balls_gate"] == pytest.approx(1.0)
    assert player_a["bowling_economy_component"] == pytest.approx(0.4)


def test_export_configurable_scoring_workbook_creates_formula_driven_top5(tmp_path) -> None:
    scorecards_dir = tmp_path / "scorecards"
    scorecards_dir.mkdir()
    (scorecards_dir / "past_matches.json").write_text(
        json.dumps(
            [
                {
                    "match_id": 101,
                    "team_a": "Team 1",
                    "team_b": "Team 2",
                    "match_start_time": "2026-04-01T09:00:00.000Z",
                }
            ]
        ),
        encoding="utf-8",
    )
    (scorecards_dir / "101_team-1-vs-team-2.json").write_text(
        json.dumps(
            {
                "match_id": 101,
                "match_title": "Team 1 vs Team 2",
                "scorecard": [
                    {
                        "teamName": "Team 1",
                        "batting": [{"name": "Player A", "runs": 20, "balls": 10, "4s": 2, "6s": 1}],
                        "bowling": [{"name": "Player C", "overs": 2, "balls": 0, "runs": 18, "wickets": 0}],
                    },
                    {
                        "teamName": "Team 2",
                        "batting": [{"name": "Player C", "runs": 12, "balls": 10, "4s": 1, "6s": 0}],
                        "bowling": [{"name": "Player A", "overs": 2, "balls": 0, "runs": 12, "wickets": 2}],
                    },
                ],
            }
        ),
        encoding="utf-8",
    )

    output_path = tmp_path / "configurable_scoring.xlsx"
    export_configurable_scoring_workbook(
        scorecards_dir,
        output_path,
        default_batting_option_2={"runs": 0.75, "delta_sr": 0.25},
        default_bowling_option_2={"wickets": 0.7, "economy": 0.3},
        default_minimum_balls_faced=0,
        default_minimum_balls_bowled=0,
    )

    workbook = load_workbook(output_path, data_only=False)
    assert output_path.exists()
    assert workbook.sheetnames == [
        "Inputs",
        "BattingInnings",
        "BowlingInnings",
        "BattingSummary",
        "BowlingSummary",
        "CombinedScores",
        "Top5",
    ]
    assert workbook["Inputs"]["B3"].value == 0.25
    assert workbook["Inputs"]["B4"].value == 0
    assert workbook["Inputs"]["B5"].value == 0.7
    assert workbook["Inputs"]["B6"].value == 0.3
    assert workbook["Inputs"]["B7"].value == 0
    assert workbook["Inputs"]["A9"].value.startswith("$$V_{bat,2}")
    assert workbook["Inputs"]["A12"].value.startswith("$$V_{bowl,2}")
    assert workbook["BattingInnings"]["G2"].value == "=IF(F2=0,0,E2/F2*100)"
    assert workbook["BattingInnings"]["N2"].value.startswith("=Inputs!$B$2*L2+Inputs!$B$3*M2")
    assert workbook["CombinedScores"]["C2"].value.startswith("=SUMIFS(BattingSummary!$F$2:")
    assert workbook["CombinedScores"]["L2"].value.startswith("=RANK.EQ(")
    assert workbook["Top5"]["B2"].value.startswith("=IFERROR(INDEX(CombinedScores!")
